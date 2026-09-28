import pytest
import numpy as np
import pandas as pd
from nifty_rl_trader.env.nifty_futures_env import NiftyFuturesEnv
from nifty_rl_trader.features.technical_features import load_config
from nifty_rl_trader.paper_trade.excel_logger import ExcelLogger
import os

@pytest.fixture
def dummy_config():
    config = load_config()
    # Override for testing deterministic behaviour
    config['settings']['env']['brokerage_per_order'] = 20
    config['settings']['env']['slippage'] = 1.0
    config['settings']['env']['lot_size'] = 25
    return config

@pytest.fixture
def dummy_df():
    dates = pd.date_range("2024-01-01 09:15:00", periods=5, freq="1min")
    df = pd.DataFrame({
        'date': dates,
        'time': dates.time,
        'open': [1000, 1010, 1020, 1030, 1040],
        'high': [1005, 1015, 1025, 1035, 1045],
        'low': [995, 1005, 1015, 1025, 1035],
        'close': [1010, 1020, 1030, 1040, 1050],
        'volume': [100, 200, 150, 300, 250],
        'log_return_1m': [0.01]*5,
        'atr_14': [0.005]*5,
        'rsi_14': [0.6]*5,
        'macd_hist': [0.001]*5,
        'bb_width': [0.02]*5,
        'vwap_div': [0.003]*5,
        'time_sin': [0.1]*5,
        'time_cos': [0.9]*5
    })
    return df

def test_observation_space(dummy_df, dummy_config):
    env = NiftyFuturesEnv(dummy_df, dummy_config)
    obs, info = env.reset()

    assert not np.isnan(obs).any(), "Observation contains NaNs"
    assert obs.shape == env.observation_space.shape, "Observation shape mismatch"

def test_transaction_costs(dummy_df, dummy_config):
    env = NiftyFuturesEnv(dummy_df, dummy_config)
    env.reset()

    # Take a Long position
    obs, reward, done, trunc, info = env.step(1)

    # Exec price is next open = 1010, slippage = 1.0 => entry = 1011
    # Cost = Brokerage (20) + Exch + SEBI + GST. Sell STT is 0 for buy.
    expected_entry = 1011.0
    turnover = expected_entry * 25
    expected_cost = 20 + (turnover * dummy_config['settings']['env']['exchange_charges']) + \
                    (turnover * dummy_config['settings']['env']['sebi_charges']) + \
                    ((20 + turnover * dummy_config['settings']['env']['exchange_charges']) * dummy_config['settings']['env']['gst_rate'])

    assert env.position == 1
    assert env.entry_price == expected_entry
    assert np.isclose(env.balance, env.initial_balance - expected_cost), f"Expected balance {env.initial_balance - expected_cost}, got {env.balance}"

def test_action_switches(dummy_df, dummy_config):
    env = NiftyFuturesEnv(dummy_df, dummy_config)
    env.reset()

    # 1. Take a Long position (Action 1)
    env.step(1)

    # Capture balance before switch
    balance_before_switch = env.balance
    entry_price_long = env.entry_price

    # 2. Switch directly to Short (Action 2)
    obs, reward, done, trunc, info = env.step(2)

    assert env.position == -1

    # Next open is 1020, slippage is 1.0
    # Long exit: 1020 - 1.0 = 1019.0
    # Short entry: 1020 - 1.0 = 1019.0

    # Long close cost calculation
    sell_turnover = 1019.0 * 25
    sell_cost = 20 + (sell_turnover * dummy_config['settings']['env']['exchange_charges']) + \
                (sell_turnover * dummy_config['settings']['env']['sebi_charges']) + \
                (sell_turnover * dummy_config['settings']['env']['stt_rate_sell']) + \
                ((20 + sell_turnover * dummy_config['settings']['env']['exchange_charges']) * dummy_config['settings']['env']['gst_rate'])

    gross_pnl_long = (1019.0 - entry_price_long) * 25
    net_pnl_long = gross_pnl_long - sell_cost

    # Short open cost calculation (Opening short is a 'sell', so it includes STT)
    sell_short_turnover = 1019.0 * 25
    sell_short_cost = 20 + (sell_short_turnover * dummy_config['settings']['env']['exchange_charges']) + \
               (sell_short_turnover * dummy_config['settings']['env']['sebi_charges']) + \
               (sell_short_turnover * dummy_config['settings']['env']['stt_rate_sell']) + \
               ((20 + sell_short_turnover * dummy_config['settings']['env']['exchange_charges']) * dummy_config['settings']['env']['gst_rate'])

    expected_balance = balance_before_switch + net_pnl_long - sell_short_cost

    assert np.isclose(env.balance, expected_balance), f"Expected balance {expected_balance}, got {env.balance}"

def test_excel_logger():
    filepath = "tests/test_log.xlsx"
    if os.path.exists(filepath):
        os.remove(filepath)

    logger = ExcelLogger(filepath)

    logger.log_trade(
        "09:15", "09:20", "LONG", 1000, 1010, 25, 250, 50, 200, "Signal", 200
    )
    logger.log_daily_summary("2024-01-01", 1, 1.0, 0, 200, 100200)

    assert os.path.exists(filepath)

    import openpyxl
    wb = openpyxl.load_workbook(filepath)
    assert "Trade_Log" in wb.sheetnames
    assert "Daily_Summary" in wb.sheetnames

    sheet = wb["Trade_Log"]
    assert sheet.max_row == 2 # Header + 1 data row
    assert sheet.cell(row=2, column=10).value == 200 # Net PnL is 200

    os.remove(filepath)
