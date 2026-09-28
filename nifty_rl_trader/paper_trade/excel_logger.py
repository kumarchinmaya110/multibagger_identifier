import os
import threading
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, numbers

class ExcelLogger:
    def __init__(self, filepath="nifty_rl_trader/trades_log.xlsx"):
        self.filepath = filepath
        self.lock = threading.Lock()

        # Setup workbook
        if not os.path.exists(self.filepath):
            self.wb = openpyxl.Workbook()
            self.trade_sheet = self.wb.active
            self.trade_sheet.title = "Trade_Log"

            # Trade Log Headers
            headers = [
                "Trade_ID", "Entry_Time", "Exit_Time", "Type", "Entry_Price",
                "Exit_Price", "Quantity", "Gross_PnL", "Charges", "Net_PnL",
                "Exit_Reason", "Cumulative_PnL"
            ]
            self.trade_sheet.append(headers)

            # Formatting headers
            for col in range(1, len(headers) + 1):
                cell = self.trade_sheet.cell(row=1, column=col)
                cell.font = Font(bold=True)

            # Summary Sheet
            self.summary_sheet = self.wb.create_sheet(title="Daily_Summary")
            summary_headers = [
                "Date", "Total_Trades", "Win_Rate_%", "Max_Drawdown_Pts",
                "Net_Profit_INR", "Ending_Capital"
            ]
            self.summary_sheet.append(summary_headers)
            for col in range(1, len(summary_headers) + 1):
                cell = self.summary_sheet.cell(row=1, column=col)
                cell.font = Font(bold=True)

            self.wb.save(self.filepath)
        else:
            self.wb = openpyxl.load_workbook(self.filepath)
            self.trade_sheet = self.wb["Trade_Log"]
            self.summary_sheet = self.wb["Daily_Summary"]

        self.trade_id_counter = self.trade_sheet.max_row # Start from last row

    def log_trade(self, entry_time, exit_time, trade_type, entry_price, exit_price, qty, gross_pnl, charges, net_pnl, exit_reason, cum_pnl):
        with self.lock:
            self.trade_id_counter += 1
            row_data = [
                self.trade_id_counter - 1, # Trade_ID
                entry_time, exit_time, trade_type,
                entry_price, exit_price, qty,
                gross_pnl, charges, net_pnl,
                exit_reason, cum_pnl
            ]

            self.trade_sheet.append(row_data)

            # Format currency and colors
            current_row = self.trade_sheet.max_row

            # Currency formats
            for col_idx in [5, 6, 8, 9, 10, 12]:
                self.trade_sheet.cell(row=current_row, column=col_idx).number_format = '₹#,##0.00'

            # Color coding Net PnL
            net_pnl_cell = self.trade_sheet.cell(row=current_row, column=10)
            if net_pnl > 0:
                net_pnl_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Green
            elif net_pnl < 0:
                net_pnl_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Red

            self.wb.save(self.filepath)

    def log_daily_summary(self, date_str, total_trades, win_rate, max_dd_pts, net_profit, ending_capital):
        with self.lock:
            row_data = [
                date_str, total_trades, win_rate, max_dd_pts, net_profit, ending_capital
            ]
            self.summary_sheet.append(row_data)

            current_row = self.summary_sheet.max_row

            # Formatting
            self.summary_sheet.cell(row=current_row, column=3).number_format = '0.00%'
            for col_idx in [5, 6]:
                self.summary_sheet.cell(row=current_row, column=col_idx).number_format = '₹#,##0.00'

            profit_cell = self.summary_sheet.cell(row=current_row, column=5)
            if net_profit > 0:
                profit_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif net_profit < 0:
                profit_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            self.wb.save(self.filepath)
